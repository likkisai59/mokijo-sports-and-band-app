"""Generate Mokijo Sports + Band Connect skills and workflow Word document."""
from docx import Document
from docx.shared import Inches, Pt, RGBColor, Cm, Emu
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml.ns import qn, nsdecls
from docx.oxml import parse_xml, OxmlElement
from pathlib import Path

OUT = Path(__file__).parent / "Mokijo-Sports-Skills-and-Workflow-Guide.docx"

NAVY = RGBColor(0x0F, 0x17, 0x2A)
TEAL = RGBColor(0x0F, 0x76, 0x6E)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
DARK = RGBColor(0x11, 0x18, 0x27)
MUTED = RGBColor(0x47, 0x55, 0x69)

# Role colors (hex without #)
C_PLAYER = "2563EB"
C_CLUB = "16A34A"
C_MEMBER = "0891B2"
C_TRAINER = "EA580C"
C_VENUE = "7C3AED"
C_SUPER = "DC2626"
C_MUKIJO = "CA8A04"
C_MONEY = "059669"
C_WAIT = "D97706"
C_HEADER = "0F766E"
C_NAVY = "0F172A"
C_LIGHT = "F0FDFA"
C_ROW_ALT = "F8FAFC"
C_SPORTS = "0369A1"
C_BAND = "7C3AED"


def shade(cell, hex_color):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_color)
    tcPr.append(shd)


def set_cell_borders(cell, color="CBD5E1", sz="4"):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcBorders = OxmlElement("w:tcBorders")
    for edge in ("top", "left", "bottom", "right"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), sz)
        el.set(qn("w:space"), "0")
        el.set(qn("w:color"), color)
        tcBorders.append(el)
    tcPr.append(tcBorders)


def set_run_font(run, name="Calibri", size=11, bold=False, color=None, italic=False):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    if color is not None:
        run.font.color.rgb = color


def add_page_border(section):
    sectPr = section._sectPr
    pgBorders = OxmlElement("w:pgBorders")
    pgBorders.set(qn("w:offsetFrom"), "page")
    for edge in ("top", "left", "bottom", "right"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), "12")
        el.set(qn("w:space"), "14")
        el.set(qn("w:color"), "0F766E")
        pgBorders.append(el)
    sectPr.append(pgBorders)


def para(doc, text, size=11, bold=False, color=DARK, space_after=8, space_before=0, align=None, italic=False):
    p = doc.add_paragraph()
    if align:
        p.alignment = align
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.space_before = Pt(space_before)
    p.paragraph_format.line_spacing = 1.15
    run = p.add_run(text)
    set_run_font(run, size=size, bold=bold, color=color, italic=italic)
    return p


def heading(doc, text, level=1):
    p = doc.add_heading(text, level=level)
    for run in p.runs:
        if level == 1:
            run.font.color.rgb = TEAL
            run.font.size = Pt(22)
        elif level == 2:
            run.font.color.rgb = RGBColor(0x0E, 0x74, 0x90)
            run.font.size = Pt(16)
        else:
            run.font.color.rgb = RGBColor(0x1E, 0x3A, 0x8A)
            run.font.size = Pt(13)
    p.paragraph_format.space_before = Pt(16 if level == 1 else 12)
    p.paragraph_format.space_after = Pt(8)
    return p


def bullet(doc, text, bold_prefix=None, size=11):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.left_indent = Inches(0.35)
    if bold_prefix:
        r1 = p.add_run(bold_prefix)
        set_run_font(r1, size=size, bold=True, color=NAVY)
        r2 = p.add_run(text)
        set_run_font(r2, size=size, color=DARK)
    else:
        r = p.add_run(text)
        set_run_font(r, size=size, color=DARK)
    return p


def numbered(doc, text):
    p = doc.add_paragraph(style="List Number")
    p.paragraph_format.space_after = Pt(4)
    r = p.add_run(text)
    set_run_font(r, size=11, color=DARK)
    return p


def callout(doc, title, body, fill="ECFDF5", title_color=TEAL):
    table = doc.add_table(rows=1, cols=1)
    table.autofit = True
    cell = table.cell(0, 0)
    shade(cell, fill)
    set_cell_borders(cell, "0F766E", "12")
    p1 = cell.paragraphs[0]
    r1 = p1.add_run(title)
    set_run_font(r1, size=12, bold=True, color=title_color)
    p2 = cell.add_paragraph()
    r2 = p2.add_run(body)
    set_run_font(r2, size=11, color=DARK)
    doc.add_paragraph()


def set_col_widths(table, widths):
    for row in table.rows:
        for i, w in enumerate(widths):
            row.cells[i].width = Inches(w)


def fill_header_row(row, headers, fill=C_HEADER, font_color=WHITE):
    for i, h in enumerate(headers):
        cell = row.cells[i]
        shade(cell, fill)
        set_cell_borders(cell, fill, "4")
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        # clear existing
        p.clear() if hasattr(p, "clear") else None
        if cell.paragraphs[0].runs:
            cell.paragraphs[0].runs[0].text = h
            set_run_font(cell.paragraphs[0].runs[0], size=11, bold=True, color=font_color)
        else:
            run = p.add_run(h)
            set_run_font(run, size=11, bold=True, color=font_color)


def add_table(doc, headers, rows, header_fill=C_HEADER, col_widths=None, first_col_fill=None):
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True

    for i, h in enumerate(headers):
        cell = table.rows[0].cells[i]
        shade(cell, header_fill)
        set_cell_borders(cell, header_fill, "4")
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        p = cell.paragraphs[0]
        run = p.add_run(h)
        set_run_font(run, size=10.5, bold=True, color=WHITE)

    for r_i, row in enumerate(rows):
        bg = C_ROW_ALT if r_i % 2 else "FFFFFF"
        for c_i, val in enumerate(row):
            cell = table.rows[r_i + 1].cells[c_i]
            fill = first_col_fill if (first_col_fill and c_i == 0) else bg
            shade(cell, fill)
            set_cell_borders(cell, "CBD5E1", "4")
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
            p = cell.paragraphs[0]
            is_first = c_i == 0
            run = p.add_run(str(val))
            set_run_font(
                run,
                size=10,
                bold=is_first,
                color=WHITE if (first_col_fill and c_i == 0) else DARK,
            )

    if col_widths:
        set_col_widths(table, col_widths)
    doc.add_paragraph()
    return table


def pin_strip(doc, steps, colors):
    """A single-row colorful pin-to-pin flow."""
    table = doc.add_table(rows=1, cols=len(steps))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, (step, color) in enumerate(zip(steps, colors)):
        cell = table.rows[0].cells[i]
        shade(cell, color)
        set_cell_borders(cell, color, "8")
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(step)
        set_run_font(run, size=9, bold=True, color=WHITE)
    doc.add_paragraph()


def flow_steps_table(doc, steps):
    """steps: list of (pin, who, action, next/result, color_hex)"""
    headers = ["Pin", "Who", "What happens", "Next result"]
    rows = [(a, b, c, d) for a, b, c, d, _ in steps]
    table = doc.add_table(rows=1 + len(rows), cols=4)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER

    for i, h in enumerate(headers):
        cell = table.rows[0].cells[i]
        shade(cell, C_NAVY)
        set_cell_borders(cell, C_NAVY, "4")
        p = cell.paragraphs[0]
        run = p.add_run(h)
        set_run_font(run, size=10.5, bold=True, color=WHITE)

    for r_i, (pin, who, action, nxt, color) in enumerate(steps):
        vals = [pin, who, action, nxt]
        for c_i, val in enumerate(vals):
            cell = table.rows[r_i + 1].cells[c_i]
            if c_i == 0:
                shade(cell, color)
                fc = WHITE
                bold = True
            else:
                shade(cell, "FFFFFF" if r_i % 2 else "F8FAFC")
                fc = DARK
                bold = c_i == 1
            set_cell_borders(cell, "CBD5E1", "4")
            p = cell.paragraphs[0]
            run = p.add_run(str(val))
            set_run_font(run, size=10, bold=bold, color=fc)
    set_col_widths(table, [0.7, 1.5, 2.8, 2.3])
    doc.add_paragraph()


def add_footer(section, text):
    footer = section.footer
    footer.is_linked_to_previous = False
    p = footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(text)
    set_run_font(run, size=8, color=MUTED, italic=True)

    # page number
    p2 = footer.add_paragraph()
    p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run2 = p2.add_run("Page ")
    set_run_font(run2, size=8, color=MUTED)
    fldChar1 = OxmlElement("w:fldChar")
    fldChar1.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    fldChar2 = OxmlElement("w:fldChar")
    fldChar2.set(qn("w:fldCharType"), "end")
    r3 = p2.add_run()
    r3._r.append(fldChar1)
    r3._r.append(instr)
    r3._r.append(fldChar2)
    set_run_font(r3, size=8, color=MUTED)


def build():
    doc = Document()

    # page setup
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.left_margin = Inches(0.75)
    section.right_margin = Inches(0.75)
    section.top_margin = Inches(0.7)
    section.bottom_margin = Inches(0.75)
    add_page_border(section)
    add_footer(section, "Mokijo Sports & Band Connect  |  Internal product & skills guide  |  Confidential")

    # ---------- COVER ----------
    para(doc, "", space_after=24)
    banner = doc.add_table(rows=3, cols=1)
    c0 = banner.cell(0, 0)
    shade(c0, C_NAVY)
    set_cell_borders(c0, C_NAVY, "4")
    p = c0.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("\nMOKIJO")
    set_run_font(r, size=36, bold=True, color=WHITE)
    p2 = c0.add_paragraph()
    p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r2 = p2.add_run("SPORTS  +  BAND CONNECT")
    set_run_font(r2, size=16, bold=True, color=RGBColor(0x5E, 0xEA, 0xD4))

    c1 = banner.cell(1, 0)
    shade(c1, C_HEADER)
    set_cell_borders(c1, C_HEADER, "4")
    p = c1.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("Senior Skills  •  Where to Use Them  •  Pin-to-Pin Workflows")
    set_run_font(r, size=14, bold=True, color=WHITE)

    c2 = banner.cell(2, 0)
    shade(c2, "134E4A")
    set_cell_borders(c2, "134E4A", "4")
    p = c2.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("Team handbook so any teammate can see who does what, which skill to apply, and how each product flow works — from register to pay to launch.\n")
    set_run_font(r, size=11, color=WHITE)

    para(doc, "", space_after=10)
    add_table(
        doc,
        ["Item", "Detail"],
        [
            ["Product", "Mokijo Sports (clubs, venues, trainers, games) + Band Connect (artists, venues, bookings)"],
            ["Audience", "Frontend, backend, and product teammates"],
            ["Stack", "React 19 + Vite + Tailwind 4  |  FastAPI + PostgreSQL + Razorpay"],
            ["Rule", "Sports team never edits Band files. Band team never edits Sports files."],
            ["Date", "13 August 2026"],
        ],
        header_fill=C_NAVY,
        col_widths=[1.6, 5.7],
    )

    heading(doc, "How to read this document", 2)
    bullet(doc, " Colors = people. Blue player, green club admin, orange trainer, purple venue owner, red super admin, gold Mukijo admin.")
    bullet(doc, " Pin-to-pin means every step is numbered: who acts → what they do → what status changes → what happens next.")
    bullet(doc, " Fundraising is the template: create campaign → share → donate → pay → raised total updates. Other flows follow the same style.")
    bullet(doc, " Skills tables tell you which senior skill to use in which area of the codebase.")

    # ---------- TOC ----------
    heading(doc, "Contents", 1)
    toc = [
        "1.  Product at a glance",
        "2.  Skills to finish frontend and backend fast",
        "3.  Senior developer skills to launch (Sports)",
        "4.  Senior developer skills to launch (Band Connect)",
        "5.  How and where to use each skill",
        "6.  Senior frontend developer (10+ years) skill set",
        "7.  Mokijo Sports — colorful pin-to-pin workflows",
        "     7.1 Who does what",
        "     7.2 Register → approve → dashboard",
        "     7.3 Venue booking (money flow)",
        "     7.4 Fundraising (full example)",
        "     7.5 Trainer enroll",
        "     7.6 Pickup games lobby",
        "     7.7 Club ops and live scoreboard",
        "     7.8 Venue verification (trust loop)",
        "8.  Launch sequence and daily usage",
        "9.  Folder map (Sports vs Band — do not mix)",
    ]
    for item in toc:
        para(doc, item, size=12, color=RGBColor(0x0F, 0x17, 0x2A), space_after=4)

    # ============================================================
    heading(doc, "1. Product at a glance", 1)
    para(
        doc,
        "This repository hosts two products that share hosting, auth patterns, and Razorpay — but they are separate businesses. Finish each as its own vertical: API → database rule → UI page → role guard.",
    )
    add_table(
        doc,
        ["Product", "Who it is for", "Core loop to launch"],
        [
            [
                "Mokijo Sports",
                "Players, club admins, members, trainers, venue owners, Super Admin, Mukijo Admin",
                "Register → (approval) → discover → book/enroll/donate → Razorpay → dashboard",
            ],
            [
                "Band Connect",
                "Clients, artists, band venues",
                "Register → profile → discover artist/venue → book → message → pay → review",
            ],
        ],
        header_fill=C_SPORTS,
        col_widths=[1.6, 2.6, 3.1],
    )

    heading(doc, "Color legend (use everywhere below)", 2)
    add_table(
        doc,
        ["Color", "Person / system"],
        [
            ["Blue", "Player / User — books venues, joins games, enrolls trainings, donates"],
            ["Green", "Club Admin — groups, events, matches, fundraising, club fees"],
            ["Cyan / Teal", "Club Member — RSVP, attendance, pay fees"],
            ["Orange", "Trainer — create trainings, see registrations"],
            ["Purple", "Venue Owner — slots, bookings, verification, payouts"],
            ["Red", "Super Admin — approve clubs and venue owners"],
            ["Gold / Yellow", "Mukijo Admin — venue document / GPS verification"],
            ["Emerald", "Money path — Razorpay order → pay → verify → paid"],
            ["Amber", "Waiting / pending approval / hold / waitlist"],
        ],
        header_fill=C_NAVY,
        first_col_fill=C_HEADER,
        col_widths=[1.8, 5.5],
    )

    # ============================================================
    heading(doc, "2. Skills to finish frontend and backend fast", 1)
    para(doc, "Must-have skills that unblock shipping. Learn or staff these first.")

    heading(doc, "Must-have (ship blockers)", 2)
    add_table(
        doc,
        ["Area", "Skill", "Description"],
        [
            ["Frontend", "React 19 + Vite + React Router", "Build Sports/Band pages, layouts, and role-based navigation."],
            ["Frontend", "Tailwind CSS 4", "Keep UI consistent and responsive without a new design system."],
            ["Frontend data", "Axios + TanStack Query", "Load, cache, and refresh API data; handle loading and errors."],
            ["Frontend state", "Zustand", "Store login/session/UI state only — not server lists."],
            ["Forms", "React Hook Form + Zod", "Reliable signup, profile, booking, and donate forms with clear errors."],
            ["Backend", "Python + FastAPI + Pydantic v2", "REST APIs with strict request/response contracts."],
            ["Database", "SQLAlchemy 2 + PostgreSQL + Alembic", "Models, queries, and safe production migrations."],
            ["Auth", "JWT + bcrypt + role checks", "Login/register for each role, enforced on the API not only the UI."],
        ],
        header_fill=C_SPORTS,
        col_widths=[1.5, 2.2, 3.6],
    )

    heading(doc, "Important for “done” features", 2)
    add_table(
        doc,
        ["Skill", "Description"],
        [
            ["Bookings / venues / trainers CRUD", "Create, list, update, cancel — with matching dashboard screens."],
            ["Razorpay payments", "Order create → checkout → verify → webhook/idempotency so money and booking stay in sync."],
            ["File / image upload", "Multipart + validation + S3/boto3 for profiles and venues."],
            ["Background jobs", "Celery/Redis or APScheduler for hold expiry, emails, retries."],
            ["API contracts", "Frontend pages wired to real FastAPI routes — no mock-only launch."],
        ],
        header_fill=C_HEADER,
        col_widths=[2.6, 4.7],
    )

    heading(doc, "Speed rules for this repo", 2)
    numbered(doc, "Use scope isolation: Sports work stays in Sports folders; Band stays in Band folders.")
    numbered(doc, "Work only in permitted paths (see Section 9).")
    numbered(doc, "Finish vertical slices: API → schema → UI page → auth guard, one feature at a time.")
    numbered(doc, "Do not block the web MVP on Flutter/mobile.")

    callout(
        doc,
        "Fastest path",
        "Pick incomplete flows (register → dashboard → venue/trainer booking or Band discovery → book → pay). Close backend endpoints first, then wire existing React pages.",
    )

    # ============================================================
    heading(doc, "3. Senior developer skills to launch (Mokijo Sports)", 1)

    heading(doc, "3.1 Product and delivery ownership", 2)
    add_table(
        doc,
        ["Skill", "Description"],
        [
            ["MVP scoping", "Cut to launchable core: auth, roles, venues/trainers, bookings, payments. Defer polish/mobile."],
            ["Feature slicing", "Ship end-to-end vertical slices instead of unfinished layers."],
            ["Release readiness", "Env configs, migrations, smoke tests, error monitoring, rollback plan."],
            ["Cross-module isolation", "Never break Band Connect routes, tables, or auth while shipping Sports."],
        ],
        header_fill=C_CLUB,
    )

    heading(doc, "3.2 Frontend (React / Vite)", 2)
    add_table(
        doc,
        ["Skill", "Description"],
        [
            ["React architecture", "Maintainable pages under frontend/app/mokijo/ and Sports dashboards."],
            ["Routing and access control", "Protect routes by role: user, trainer, venue owner, club admin, super admin."],
            ["Server-state management", "TanStack Query: caching, invalidation, loading/error, retries."],
            ["Client-state management", "Zustand for auth/session only; do not duplicate server lists."],
            ["Forms and validation", "RHF + Zod with API error mapping for signup, booking, donate, profile."],
            ["UI systems (Tailwind)", "Consistent Sports UI without breaking existing patterns."],
            ["API integration quality", "Axios auth headers, interceptors, toasts, user-safe errors."],
            ["Performance and UX polish", "Empty/loading/error states, no double-submit, slow-network handling."],
        ],
        header_fill=C_PLAYER,
    )

    heading(doc, "3.3 Backend (FastAPI / Python)", 2)
    add_table(
        doc,
        ["Skill", "Description"],
        [
            ["FastAPI API design", "Clear REST for venues, trainers, bookings, matches, events, fundraising, admin."],
            ["Pydantic contracts", "Invalid input fails early; frontend and backend stay aligned."],
            ["SQLAlchemy + PostgreSQL", "Relationships, indexes, no N+1, no double-book races."],
            ["Alembic migrations", "Safe schema evolution. Never drop Band Connect band_* tables."],
            ["Auth and authorization", "JWT and role checks on every sensitive endpoint."],
            ["Business rules", "Holds, expiry, availability, cancellations, ownership — in services, not only UI."],
            ["Payments (Razorpay)", "Order → verify → webhook; booking/payment status stay consistent."],
            ["Async / jobs", "Hold expiry, emails, retries via Celery/Redis or APScheduler."],
            ["File / media", "Secure uploads for profiles and venues."],
            ["API security basics", "Rate limit, CORS, secrets in env, least-privilege DB."],
        ],
        header_fill="1E3A8A",
    )

    heading(doc, "3.4 Data, reliability, and ops (launch blockers)", 2)
    add_table(
        doc,
        ["Skill", "Description"],
        [
            ["Environment management", "Separate local / staging / prod. No secrets in git."],
            ["Observability", "Logs, health checks, know when bookings or payments fail in prod."],
            ["Backup and recovery", "DB backups, migration rollback, payment reconciliation."],
            ["CI/CD basics", "Lint, test, build, deploy for frontend + backend."],
            ["Hosting", "FastAPI + Postgres + Redis (+ worker), HTTPS, domains, CORS."],
            ["Smoke testing", "Register → login → book → pay → dashboard. Protect Sports vs Band regressions."],
        ],
        header_fill="0F172A",
    )

    heading(doc, "3.5 Sports domain skills", 2)
    add_table(
        doc,
        ["Skill", "Description"],
        [
            ["Multi-role marketplace", "User, trainer, venue owner, club admin, member, platform admin journeys."],
            ["Scheduling and inventory", "Slots, double-booking prevention, holds/expiry."],
            ["Admin operations", "Approve/reject clubs, owners, venues so the platform can be operated."],
            ["Notifications", "Email or in-app for booking confirmations and failures."],
        ],
        header_fill=C_CLUB,
    )

    heading(doc, "Sports launch-priority order", 2)
    pin_strip(
        doc,
        ["1 Auth + roles", "2 Profiles CRUD", "3 Bookings", "4 Razorpay", "5 Admin + email", "6 Staging → Prod"],
        [C_PLAYER, C_MEMBER, C_WAIT, C_MONEY, C_SUPER, C_HEADER],
    )

    # ============================================================
    heading(doc, "4. Senior developer skills to launch (Band Connect)", 1)
    para(
        doc,
        "Band Connect is an artist/venue marketplace: discover → request booking → confirm → pay → message → review. Work only in Band folders.",
    )

    heading(doc, "4.1 Product and delivery (Band)", 2)
    add_table(
        doc,
        ["Skill", "Description"],
        [
            ["Artist marketplace MVP", "Minimum loop: discover artists/venues → request booking → confirm → pay → review."],
            ["Multi-role Band UX", "Distinct journeys for client, artist, and venue."],
            ["Scope isolation", "Only frontend/app/band/, frontend/app/artist/, backend/app/api/band/, band_models.py."],
            ["Launch checklist", "Auth, profiles, calendar, booking states, payments, notifications, messaging."],
        ],
        header_fill=C_BAND,
    )

    heading(doc, "4.2 Frontend skills (Band UI)", 2)
    add_table(
        doc,
        ["Skill", "Where it is used", "Description"],
        [
            ["Band app routing", "/band and /artist", "Login, register, dashboards, bookings, messages."],
            ["Role-based dashboards", "artist / venue / client dashboards", "Correct nav, empty states, permission-aware actions."],
            ["Discovery UX", "artists list + venues list + [id] pages", "Filters (category, location) mapped to APIs."],
            ["Booking and calendar UI", "artist/calendar, bookings pages", "Pending / confirmed / cancelled users can trust."],
            ["Messaging UI", "inbox + thread pages", "Send/receive, unread, error handling."],
            ["Reviews and earnings UI", "artist/reviews, earnings", "Wired to real APIs, not placeholders."],
            ["Forms + Zod", "register, profile, settings", "API error mapping."],
            ["Query / Zustand split", "all Band screens", "Invalidate after book / message / pay."],
            ["Media / crop UX", "artist and venue profiles", "react-easy-crop with loading/failure states."],
        ],
        header_fill=C_BAND,
        col_widths=[2.0, 2.2, 3.1],
    )

    heading(doc, "4.3 Backend skills (Band API)", 2)
    add_table(
        doc,
        ["Skill", "Description"],
        [
            ["Band FastAPI modules", "Auth, artists, venues, bookings, categories, locations, reviews, earnings, payments, notifications, messaging, settings."],
            ["Band domain modeling", "Accounts, artist profiles, venues, bookings, reviews, transactions, payment orders, conversations."],
            ["Auth for Band roles", "JWT/refresh; artist vs client vs venue on every endpoint."],
            ["Booking lifecycle", "Request → accept/reject → confirm → complete/cancel with ownership checks."],
            ["Availability", "Prevent double-booking artists/venues; validate date/time windows."],
            ["Locations and categories", "Country/state/city/area + category taxonomy for search."],
            ["Razorpay Band", "Payment order create/verify/webhook + booking status consistency."],
            ["Earnings", "Aggregate successful transactions for artist/venue reporting."],
            ["Reviews", "Post-booking only; rating aggregation on profiles."],
            ["Messaging backend", "Participant authorization, unread counts, pagination."],
            ["Notifications", "Booking and message events; optional SendGrid email."],
            ["Storage", "Secure media for artist/venue assets."],
        ],
        header_fill="5B21B6",
    )

    heading(doc, "Band Connect launch-priority order", 2)
    pin_strip(
        doc,
        ["1 Auth", "2 Profiles", "3 Discovery", "4 Book + calendar", "5 Pay", "6 Chat + notify", "7 Reviews + earnings", "8 Staging"],
        [C_BAND, C_VENUE, C_PLAYER, C_WAIT, C_MONEY, C_MEMBER, C_CLUB, C_NAVY],
    )

    # ============================================================
    heading(doc, "5. How and where to use each skill", 1)
    para(doc, "Simple rule: (1) pick Sports or Band, (2) pick the area, (3) apply only that area’s skills, (4) stay in the correct folders.")

    heading(doc, "5.1 Module first", 2)
    add_table(
        doc,
        ["If you work on…", "Use skills for…", "Work only in…", "Do not touch…"],
        [
            ["Mokijo Sports", "Sports skills", "frontend/app/mokijo/, trainings, venues, bookings, dashboards + backend mokijo APIs", "Band folders"],
            ["Band Connect", "Band skills", "frontend/app/band/, frontend/app/artist/ + backend/app/api/band/", "Sports folders"],
            ["Shared files", "Additive shared-platform only", "main.py, routers, global CSS — append only", "Other module’s routes"],
        ],
        header_fill=C_NAVY,
        col_widths=[1.6, 1.6, 2.5, 1.6],
    )

    heading(doc, "A) Auth and registration", 2)
    para(doc, "Where: Sports mokijo/register-* and login pages, components/auth/. Band: band/login, band/register. Backend: Sports auth APIs and backend/app/api/band/auth/.")
    add_table(
        doc,
        ["Skills to use", "How to use"],
        [
            ["JWT auth and role permissions", "Register → login → store token → redirect by role → guard dashboards."],
            ["Forms + Zod", "Field-level errors + map API 400/401 messages."],
            ["Protected routing", "Wrong role cannot open another role’s dashboard."],
            ["Zustand session", "Auth user only; refetch profile from API when needed."],
        ],
        header_fill=C_PLAYER,
    )

    heading(doc, "B) Profiles and onboarding", 2)
    para(doc, "Where: Sports trainer/venue/user profiles. Band: artist/profile, settings, venue profile APIs.")
    add_table(
        doc,
        ["Skills to use", "How to use"],
        [
            ["React forms + image upload/crop", "Complete profile before booking/discovery is fully usable."],
            ["Pydantic + SQLAlchemy", "Validate and persist profile fields."],
            ["Storage (S3/multipart)", "Safe image types and size limits."],
            ["Role-based fields", "Trainer vs venue vs artist see different required fields."],
        ],
        header_fill=C_TRAINER,
    )

    heading(doc, "C) Discovery / listing", 2)
    para(doc, "Where: Sports venues/trainers/activities. Band: band/artists, band/venues, detail [id] pages.")
    add_table(
        doc,
        ["Skills to use", "How to use"],
        [
            ["API list/filter/pagination", "Search → open detail → CTA to book."],
            ["TanStack Query caching", "Do not refetch on every tab blink; invalidate after booking."],
            ["Categories + locations (Band)", "Filters must match backend taxonomy."],
            ["Responsive listing UI", "Cards work on mobile and desktop."],
        ],
        header_fill=C_MEMBER,
    )

    heading(doc, "D) Booking and calendar (core product)", 2)
    para(doc, "Where: Sports bookings + venue/trainer availability. Band: band/*/bookings, artist/calendar, backend bookings.")
    add_table(
        doc,
        ["Skills to use", "How to use"],
        [
            ["Booking lifecycle", "Request → accept → confirm → cancel with visible status."],
            ["Availability / anti-double-book", "Enforce in the API. Never trust UI-only checks."],
            ["Backend business rules", "Ownership, hold windows, slot state."],
            ["Query invalidation", "After status change, refresh calendar and lists."],
        ],
        header_fill=C_WAIT,
    )

    heading(doc, "E) Payments", 2)
    para(doc, "Where: Band backend/app/api/band/payments/. Sports venue booking, games, trainings, club fees, fundraising.")
    add_table(
        doc,
        ["Skills to use", "How to use"],
        [
            ["Razorpay order/verify/webhook", "Pay success must update booking + ledger in one reliable flow."],
            ["Idempotent payments", "Retries must not confirm twice."],
            ["Status sync", "UI shows pending / paid / failed / refunded from server truth."],
            ["Env keys", "Secret key server-only; frontend uses publishable/key id only."],
        ],
        header_fill=C_MONEY,
    )

    heading(doc, "F) Messaging and notifications (Band-heavy; Sports has messages too)", 2)
    add_table(
        doc,
        ["Skills to use", "How to use"],
        [
            ["Conversation authorization", "Only participants can read a thread."],
            ["Unread / pagination", "Inbox badge and load older messages."],
            ["Notification create + mark-read", "Booking/message events → notification → badge."],
            ["Optional email (SendGrid)", "Confirmations and failures."],
        ],
        header_fill=C_MEMBER,
    )

    heading(doc, "G) Reviews, earnings, analytics (Band) + fundraising (Sports)", 2)
    add_table(
        doc,
        ["Skills to use", "How to use"],
        [
            ["Post-booking review rules", "Only after a completed booking."],
            ["Earnings aggregation", "From successful transactions only."],
            ["Fundraising campaign + donate", "Create → publish → donate → pay → add to raised total (see Section 7.4)."],
            ["Empty/error states", "No blank dashboards when there is zero data."],
        ],
        header_fill=C_CLUB,
    )

    heading(doc, "H) Admin / ops", 2)
    add_table(
        doc,
        ["Skills to use", "How to use"],
        [
            ["Admin authorization", "Super Admin vs Mukijo Admin vs Club Admin are different."],
            ["Taxonomy management", "Categories and locations (Band); groups/members (Sports)."],
            ["Venue verification", "Docs + GPS review queue."],
            ["Disputes / refunds runbook", "Know how to reverse a paid booking safely."],
        ],
        header_fill=C_SUPER,
    )

    heading(doc, "I) Launch / production", 2)
    add_table(
        doc,
        ["Skills to use", "How to use"],
        [
            ["Env/secrets", "Staging first, then production cutover."],
            ["Alembic (module-safe)", "Sports migrations never touch band_*."],
            ["Smoke tests", "Critical path per product before go-live."],
            ["Logging / monitoring", "Watch payment verify and booking state changes."],
        ],
        header_fill=C_NAVY,
    )

    heading(doc, "Quick “what skill now?” cheat sheet", 2)
    add_table(
        doc,
        ["You’re stuck on…", "Use these skills"],
        [
            ["Login/register broken", "Auth JWT, forms/Zod, protected routes"],
            ["Dashboard empty / wrong role", "Role permissions, TanStack Query, routing"],
            ["Cannot find artists/venues", "Discovery filters, locations/categories APIs"],
            ["Double booking / wrong status", "Booking lifecycle + backend business rules"],
            ["Money taken but booking pending", "Razorpay + idempotency + status sync"],
            ["Chat not updating", "Messaging auth + query invalidation"],
            ["Deploy fails / prod bugs", "Env, migrations, smoke tests, observability"],
            ["Accidentally editing other module", "Scope isolation — stop and revert that path"],
        ],
        header_fill=C_NAVY,
        col_widths=[2.8, 4.5],
    )

    # ============================================================
    heading(doc, "6. Senior frontend developer (10+ years) skill set", 1)
    para(
        doc,
        "A senior frontend engineer at this level does not only make pages look good. They own end-to-end frontend quality for one module (Sports or Band), prevent rework, ship production-safe booking/auth/payment UX, and keep the other module safe.",
    )

    heading(doc, "6.1 Core craft (must-have)", 2)
    add_table(
        doc,
        ["Skill", "Description"],
        [
            ["Advanced React architecture", "Feature folders, shared UI, layouts, role-based areas that stay maintainable as features grow."],
            ["Component system design", "Reusable accessible primitives: forms, tables, dialogs, empty/loading/error — not one-off markup."],
            ["Modern React mastery", "Hooks, composition, effects; knows when not to over-optimize; predictable UI logic."],
            ["Routing and app shell", "Protected routes, role redirects, nested layouts, deep links for every dashboard."],
            ["Complex form systems", "Multi-step register, profile, booking, donate, settings with draft recovery and API error mapping."],
            ["Server vs client state", "TanStack Query for server data; Zustand for session/UI only."],
            ["API integration excellence", "Auth headers, refresh, retries, idempotent actions, user-safe errors."],
        ],
        header_fill=C_PLAYER,
    )

    heading(doc, "6.2 Product UI and UX (senior differentiator)", 2)
    add_table(
        doc,
        ["Skill", "Description"],
        [
            ["Product thinking in UI", "Turns booking/marketplace flows into simple screens. Conversion over decoration."],
            ["Information architecture", "Users always know where they are: discover → detail → book → pay → manage."],
            ["Interaction design judgment", "Loading, empty, success, failure for book, pay, message, accept/reject, donate."],
            ["Responsive UI", "Dashboards and booking flows work on desktop and mobile."],
            ["Accessibility (a11y)", "Keyboard, focus, labels, contrast — launch requirement, not polish."],
            ["Visual consistency", "Typography, spacing, Tailwind patterns so Sports/Band feel intentional."],
        ],
        header_fill=C_HEADER,
    )

    heading(doc, "6.3 Performance, quality, reliability", 2)
    add_table(
        doc,
        ["Skill", "Description"],
        [
            ["Frontend performance engineering", "Fix slow lists, extra re-renders, large bundles, request waterfalls."],
            ["Perceived performance", "Skeletons, safe optimistic UI, fast first interaction on key pages."],
            ["Defensive UI", "Partial API data, race conditions, double-submit, slow network, permission failures — no blank screens."],
            ["Testing strategy", "Critical-path flows, form validation, auth guards, booking/payment UI regression."],
            ["Observability in the UI", "Client error reporting so production bugs are diagnosable."],
        ],
        header_fill="1E3A8A",
    )

    heading(doc, "6.4 Platform and engineering leadership", 2)
    add_table(
        doc,
        ["Skill", "Description"],
        [
            ["Design-system literacy", "Radix/shadcn-style components without inconsistent one-offs."],
            ["Frontend security", "XSS, token storage, open redirects, hide privileged actions from wrong roles."],
            ["Build and delivery", "Vite/env, lint, CI frontend checks, web release smoke."],
            ["API contract collaboration", "Push for FE-friendly schemas; communicate breaking changes."],
            ["Mentorship / review", "PRs checked for architecture, a11y, state bugs, Sports vs Band scope."],
            ["Estimation and sequencing", "Auth → core flows → payments UX → polish."],
            ["Multi-module navigation", "Change only owned frontend paths; shared shells stay additive."],
        ],
        header_fill=C_NAVY,
    )

    heading(doc, "6.5 Domain skills for this product (senior FE)", 2)
    add_table(
        doc,
        ["Skill", "Description"],
        [
            ["Marketplace UI", "Discovery, trust signals, booking timelines, role dashboards."],
            ["Booking / calendar UX", "Availability, conflict feedback, pending/confirmed/cancelled."],
            ["Payments UX", "Razorpay: initiate, success, failure, pending verify, post-pay confirmation."],
            ["Messaging UI", "Inbox/thread, unread, send failure, permission-safe views."],
            ["Fundraising UX", "Campaign progress (raised/goal), donate form, success/failure, donor list for admin."],
        ],
        header_fill=C_CLUB,
    )

    heading(doc, "Where a 10+ year FE uses these skills", 2)
    add_table(
        doc,
        ["Area", "Skills to apply"],
        [
            ["Login / register", "Forms, validation, auth state, security, protected routes"],
            ["Dashboards", "IA, role UI, Query/Zustand split, empty/error states"],
            ["Listings / detail", "Performance, filtering UX, API integration"],
            ["Bookings / calendar", "Domain booking UX, defensive UI, state sync"],
            ["Payments", "Checkout UX, error recovery, status clarity"],
            ["Fundraising donate", "Amount form, Razorpay states, progress bar, receipts"],
            ["Messages / notifications", "Messaging patterns, refresh, a11y"],
            ["Shared layout / nav", "App shell, design consistency, additive shared changes"],
        ],
        header_fill=C_PLAYER,
        col_widths=[2.2, 5.1],
    )

    # ============================================================
    heading(doc, "7. Mokijo Sports — colorful pin-to-pin workflows", 1)
    callout(
        doc,
        "How to read a pin-to-pin flow",
        "Each pin is one step. Read left-to-right (or top-to-bottom). “Who” is the colored person. “What happens” is the screen or API action. “Next result” is the new status. Fundraising (7.4) is the full example of this style.",
        fill="ECFEFF",
        title_color=RGBColor(0x0E, 0x74, 0x90),
    )

    heading(doc, "7.1 Who does what (big picture)", 2)
    para(doc, "Mokijo Sports is a sports club + venue + training platform. Seven actor types:")

    roles = [
        ("Blue PLAYER / USER", "Book venues, join pickup games, enroll in trainings, donate to club campaigns, review venues.", C_PLAYER),
        ("Green CLUB ADMIN", "Manage club, groups/teams, members, events, matches, courses, fundraising, club fee payments.", C_CLUB),
        ("Cyan CLUB MEMBER", "Join via signup form, RSVP to events, attend, pay club fees, take courses.", C_MEMBER),
        ("Orange TRAINER", "Create/edit trainings, see who registered.", C_TRAINER),
        ("Purple VENUE OWNER", "List venues, set slots, approve/reject bookings, verification docs, payouts.", C_VENUE),
        ("Red SUPER ADMIN", "Approve or reject club admins and venue owners before they can operate.", C_SUPER),
        ("Gold MUKIJO ADMIN", "Review venue documents and GPS; verify, reject, request more info, or suspend.", C_MUKIJO),
    ]
    t = doc.add_table(rows=len(roles), cols=2)
    for i, (title, desc, color) in enumerate(roles):
        c0 = t.rows[i].cells[0]
        c1 = t.rows[i].cells[1]
        shade(c0, color)
        shade(c1, "F8FAFC")
        set_cell_borders(c0, color, "8")
        set_cell_borders(c1, "CBD5E1", "4")
        p0 = c0.paragraphs[0]
        r0 = p0.add_run(title)
        set_run_font(r0, size=11, bold=True, color=WHITE)
        p1 = c1.paragraphs[0]
        r1 = p1.add_run(desc)
        set_run_font(r1, size=10.5, color=DARK)
    set_col_widths(t, [2.4, 4.9])
    doc.add_paragraph()

    heading(doc, "7.2 Register → approve → dashboard (entry flow)", 2)
    para(doc, "Everyone starts at /mokijo (or /) and picks a role. Club Admin and Venue Owner wait for Super Admin. Then each role lands on its own dashboard.")
    pin_strip(
        doc,
        ["Land /mokijo", "Choose role", "Register", "Wait if needed", "Login", "Role dashboard"],
        [C_NAVY, C_HEADER, C_PLAYER, C_WAIT, C_MONEY, C_CLUB],
    )
    flow_steps_table(
        doc,
        [
            ("1", "Anyone", "Open landing /mokijo", "See role choices", C_NAVY),
            ("2a", "Player", "Register User → Login User", "User Dashboard (/user-dashboard)", C_PLAYER),
            ("2b", "Club Admin", "Register Club → status PENDING_APPROVAL", "Wait for Super Admin", C_CLUB),
            ("2c", "Super Admin", "Approve or reject club", "APPROVED → Club can log in", C_SUPER),
            ("2d", "Club Admin", "Login Club", "Club Dashboard (/dashboard)", C_CLUB),
            ("2e", "Member", "Register / signup form → Login Member", "Club dashboard (member view)", C_MEMBER),
            ("2f", "Trainer", "Register Trainer → Login Trainer", "Trainer Dashboard", C_TRAINER),
            ("2g", "Venue Owner", "Register → PENDING_APPROVAL → Super Admin", "Venue Dashboard after APPROVED", C_VENUE),
            ("2h", "Mukijo Admin", "Login Admin", "Venue verification queue", C_MUKIJO),
            ("2i", "Super Admin", "Super Admin login", "Approvals dashboard", C_SUPER),
        ],
    )
    add_table(
        doc,
        ["Role", "Register path", "Login path", "Lands on"],
        [
            ["Club Admin", "/register", "/login", "/dashboard"],
            ["Member", "/register-member or signup form", "/login-member", "/dashboard"],
            ["Player", "/register-user", "/login-user", "/user-dashboard"],
            ["Trainer", "/register-trainer", "/login-trainer", "/trainer-dashboard"],
            ["Venue Owner", "/register-venue", "/login-venue", "/venue-dashboard"],
            ["Mukijo Admin", "—", "/login-admin", "/dashboard/venue-verification"],
            ["Super Admin", "—", "/super-admin/login", "/super-admin/dashboard"],
        ],
        header_fill=C_NAVY,
        col_widths=[1.6, 2.1, 1.8, 1.8],
    )

    heading(doc, "7.3 Venue booking — pin-to-pin money flow", 2)
    para(doc, "This is the main Sports money loop. Player holds a slot, may wait for owner approval, pays on Razorpay, then the booking is confirmed.")
    pin_strip(
        doc,
        ["Open venues", "Pick date/slots", "HOLD (HELD)", "Approve?", "Razorpay pay", "Verify", "CONFIRMED"],
        [C_PLAYER, C_PLAYER, C_WAIT, C_VENUE, C_MONEY, C_MONEY, C_CLUB],
    )
    flow_steps_table(
        doc,
        [
            ("1", "Player", "Open Venues list or User Dashboard", "See available venues", C_PLAYER),
            ("2", "Player", "Open venue detail, pick date and slots", "Slots still AVAILABLE", C_PLAYER),
            ("3", "Player + API", "POST /bookings/hold", "Slots AVAILABLE → HELD (~5 min). Booking pending_payment", C_WAIT),
            ("4", "System", "Does this venue need owner approval?", "No → go to pay. Yes → pending_approval", C_WAIT),
            ("5", "Venue Owner", "Approve or reject in venue-dashboard/bookings", "Reject → slots released. Approve → player can pay", C_VENUE),
            ("6", "Player + API", "POST /bookings/razorpay/order", "Razorpay Checkout opens", C_MONEY),
            ("7", "Player", "Pay on Razorpay", "Success or fail/cancel", C_MONEY),
            ("8", "API", "POST /bookings/razorpay/verify", "Booking CONFIRMED. Payment PAID. Slots BOOKED", C_CLUB),
            ("9", "Player (optional)", "POST /bookings/{id}/cancel", "Slots freed. If paid → refunded", C_SUPER),
        ],
    )
    para(doc, "Status pin strips", size=12, bold=True, color=TEAL, space_after=6)
    para(doc, "Slots:   AVAILABLE  →  HELD  →  BOOKED   (or back to AVAILABLE on reject/cancel/expire)", size=11)
    para(doc, "Booking: reserved / pending_payment / pending_approval  →  confirmed  →  cancelled / completed", size=11)
    para(doc, "Payment: pending  →  paid  →  refunded", size=11, space_after=12)

    heading(doc, "7.4 Fundraising — full pin-to-pin example", 2)
    para(
        doc,
        "Read this flow first if you are new. Every other money flow (venue book, training enroll, game join, club fees) uses the same pattern: create something → user pays Razorpay → verify → update totals/status.",
        bold=False,
    )
    pin_strip(
        doc,
        ["Admin opens Fundraising", "Create campaign", "Publish LIVE", "Donor opens donate", "Pay Razorpay", "Verify", "Raised ₹ updates"],
        [C_CLUB, C_CLUB, C_MONEY, C_PLAYER, C_MONEY, C_MONEY, C_WAIT],
    )

    heading(doc, "Why fundraising exists", 3)
    para(
        doc,
        "Clubs need to raise money for kits, travel, grounds, or events. Mokijo lets a Club Admin publish a campaign with a rupee goal. Members, players, or anyone with the link can donate. Razorpay takes the payment. The campaign progress bar (Raised / Goal) updates only after payment is verified. Admin can see donors and the total.",
    )

    heading(doc, "Fundraising pins (detailed)", 3)
    flow_steps_table(
        doc,
        [
            ("1", "Club Admin", "Open Dashboard → Fundraising", "See existing campaigns or empty state", C_CLUB),
            ("2", "Club Admin", "Click Create Campaign", "Form: title, goal ₹, story, end date, image", C_CLUB),
            ("3", "Club Admin", "Save / Publish", "Campaign status → active (LIVE)", C_CLUB),
            ("4", "Admin or system", "Share donate link /fundraising/donate/[id]", "Donors can open the page", C_WAIT),
            ("5", "Donor (Player / Member / public)", "Open donate page, enter amount + details", "Ready to pay", C_PLAYER),
            ("6", "API", "Initiate donation + create Razorpay order", "Checkout opens. Donation = initiated", C_MONEY),
            ("7", "Donor", "Pay on Razorpay Checkout", "Success, fail, or user cancel", C_MONEY),
            ("8a", "API (success)", "Verify payment", "Donation = PAID. Add amount to campaign raised ₹", C_CLUB),
            ("8b", "System (fail)", "Payment failed or cancelled", "Donation = failed. Raised total UNCHANGED", C_SUPER),
            ("9", "System", "Is goal reached?", "No → show progress bar. Yes → campaign completed (or keep accepting if allowed)", C_WAIT),
            ("10", "Club Admin", "View donors list + total raised", "Ops complete for this campaign", C_CLUB),
        ],
    )

    para(doc, "Fundraising status pins", size=12, bold=True, color=TEAL, space_after=6)
    para(doc, "Campaign:  draft  →  active  →  completed / closed", size=11)
    para(doc, "Donation:  initiated  →  paid   or   failed", size=11)
    para(doc, "Progress:  raised = raised + donation_amount   until goal or end date", size=11, space_after=10)

    heading(doc, "Fundraising — what each person sees", 3)
    add_table(
        doc,
        ["Person", "Screen", "They can do"],
        [
            ["Club Admin (green)", "/dashboard/fundraising", "Create, publish, watch raised vs goal, see donors"],
            ["Club Admin (green)", "/dashboard/fundraising/new", "Fill campaign form"],
            ["Donor (blue)", "/dashboard/fundraising/donate/[id]", "Enter amount, pay, see success/fail"],
            ["Razorpay (emerald)", "Checkout popup / redirect", "Collect card/UPI/netbanking"],
            ["Backend", "/fundraising APIs", "Create campaign, order, verify, add to total"],
        ],
        header_fill=C_CLUB,
        col_widths=[2.0, 2.6, 2.7],
    )

    heading(doc, "Same pattern used in other money flows", 3)
    add_table(
        doc,
        ["Flow", "“Create” pin", "“Pay” pin", "“Update total / status” pin"],
        [
            ["Fundraising", "Admin publishes campaign", "Donor Razorpay", "raised ₹ += amount"],
            ["Venue booking", "Player holds slots", "Player Razorpay", "booking confirmed, slots BOOKED"],
            ["Training enroll", "Trainer publishes training", "Player Razorpay (if paid)", "registration = registered"],
            ["Pickup game", "Host creates lobby", "Joiner Razorpay after accept (if needed)", "player = confirmed"],
            ["Club fees", "Admin creates payment request", "Member Razorpay", "fee = paid (or overdue if late)"],
        ],
        header_fill=C_MONEY,
        col_widths=[1.6, 2.0, 2.0, 1.7],
    )

    heading(doc, "7.5 Trainer training enroll flow", 2)
    pin_strip(
        doc,
        ["Trainer creates", "Training OPEN", "Player discovers", "Free or paid?", "Enroll / Pay", "Trainer sees list"],
        [C_TRAINER, C_TRAINER, C_PLAYER, C_WAIT, C_MONEY, C_TRAINER],
    )
    flow_steps_table(
        doc,
        [
            ("1", "Trainer", "Create training on trainer-dashboard", "Training OPEN (or draft until published)", C_TRAINER),
            ("2", "Player", "Discover trainings on user dashboard or /trainings/[id]", "See price and details", C_PLAYER),
            ("3a", "Player", "If FREE: enroll", "Registration = registered", C_CLUB),
            ("3b", "Player", "If PAID: enroll/order → Razorpay → verify", "Registration = registered, payment = paid", C_MONEY),
            ("4", "Trainer", "Open Registrations", "See who joined", C_TRAINER),
        ],
    )
    para(doc, "Course/training pins:  draft | open | full | closed | completed.   Registration: unregistered → registered | waitlisted | cancelled.", size=10, italic=True, color=MUTED)

    heading(doc, "7.6 Pickup games lobby flow", 2)
    pin_strip(
        doc,
        ["Host creates game", "Host CONFIRMED", "Player joins", "Instant or approval?", "Pay", "CONFIRMED / waitlist"],
        [C_PLAYER, C_CLUB, C_PLAYER, C_WAIT, C_MONEY, C_MEMBER],
    )
    flow_steps_table(
        doc,
        [
            ("1", "Player (host)", "Create game (POST /games)", "Host auto-confirmed. Game = open", C_PLAYER),
            ("2", "Other player", "Join game", "Branch on join policy", C_PLAYER),
            ("3a", "Instant policy", "pending_payment → Razorpay → confirmed", "Seat taken", C_MONEY),
            ("3b", "Request approval", "pending_approval", "Wait for host", C_WAIT),
            ("4", "Host", "Accept or reject", "Accept → pending_payment. Reject → rejected", C_PLAYER),
            ("5", "Joiner", "Pay after accept (if required)", "confirmed", C_MONEY),
            ("6", "System", "Lobby full?", "Yes → extras WAITLIST. Seat free → promote waitlist", C_WAIT),
            ("7", "Host (optional)", "Cancel lobby", "Refund confirmed players", C_SUPER),
        ],
    )
    para(doc, "Game: open → full | cancelled | completed.   Player: pending_approval → pending_payment → confirmed | rejected | cancelled.   Waitlist: waiting → promoted | expired.", size=10, italic=True, color=MUTED)

    heading(doc, "7.7 Club ops and live match scoreboard", 2)
    pin_strip(
        doc,
        ["Create groups", "Add members", "Events / courses / fees", "Create match", "LIVE scoring", "Public scoreboard"],
        [C_CLUB, C_MEMBER, C_CLUB, C_SUPER, C_SUPER, C_VENUE],
    )
    flow_steps_table(
        doc,
        [
            ("1", "Club Admin", "Create groups / teams", "Club structure exists", C_CLUB),
            ("2", "Admin / Member", "Add, import, or approve signup forms (Coach / Parent / Player / Referee)", "Members in groups", C_MEMBER),
            ("3", "Club Admin", "Create events, courses, or payment requests", "Members can RSVP / register / pay", C_CLUB),
            ("4", "Member", "Respond to event: pending | accepted | declined | maybe | waitlisted", "Attendance later: present / absent / late", C_MEMBER),
            ("5", "Club Admin or Venue", "Create match (scheduled)", "Manage teams and scores", C_SUPER),
            ("6", "Scorer", "Match goes LIVE, events pushed", "WebSocket /ws/scoreboard/{id}", C_SUPER),
            ("7", "Public / parents", "Open /scoreboard/[match_id]", "Live board updates", C_VENUE),
            ("8", "System", "Match completed or cancelled", "Final score stored", C_NAVY),
        ],
    )

    heading(doc, "7.8 Venue verification — trust loop (before real bookings)", 2)
    para(doc, "A venue should be verified before it is trusted for paid bookings. Owner submits docs + GPS. Mukijo Admin reviews.")
    pin_strip(
        doc,
        ["DRAFT venue", "Docs + GPS", "PENDING", "UNDER_REVIEW", "VERIFIED / REJECT / MORE INFO", "Can take bookings"],
        [C_VENUE, C_VENUE, C_WAIT, C_MUKIJO, C_CLUB, C_MONEY],
    )
    flow_steps_table(
        doc,
        [
            ("1", "Venue Owner", "Add venue", "Status DRAFT", C_VENUE),
            ("2", "Venue Owner", "Upload documents + GPS", "Ready to submit", C_VENUE),
            ("3", "Venue Owner", "Submit for verification", "PENDING_VERIFICATION", C_WAIT),
            ("4", "Mukijo Admin", "Start review", "UNDER_REVIEW", C_MUKIJO),
            ("5a", "Mukijo Admin", "Approve", "VERIFIED → venue can take bookings", C_CLUB),
            ("5b", "Mukijo Admin", "Reject", "REJECTED", C_SUPER),
            ("5c", "Mukijo Admin", "Request more info", "MORE_INFO_REQUIRED → owner resubmits", C_WAIT),
            ("6", "Mukijo Admin (later)", "Suspend if needed", "SUSPENDED", C_SUPER),
        ],
    )

    heading(doc, "7.9 One-page map for teammates", 2)
    add_table(
        doc,
        ["Flow", "Start", "End", "Why it matters"],
        [
            ["Auth", "Role pick", "Role dashboard", "Nobody can use the product without the right login"],
            ["Venue book", "Discover venue", "Confirmed + paid (or refund)", "Core Sports revenue"],
            ["Fundraising", "Create campaign", "Raised total + donor list", "Clubs raise money for real needs"],
            ["Training", "Trainer publishes", "Player registered", "Trainer marketplace"],
            ["Games", "Host lobby", "Confirmed players / waitlist", "Pickup play with fair seats and pay"],
            ["Match", "Create match", "Public live scoreboard", "Parents and clubs watch live"],
            ["Trust", "Venue docs", "VERIFIED", "Stop fake or unsafe venues before money moves"],
        ],
        header_fill=C_NAVY,
        col_widths=[1.5, 1.6, 2.0, 2.2],
    )

    # ============================================================
    heading(doc, "8. Launch sequence and daily usage", 1)

    heading(doc, "Sports team sequence", 2)
    numbered(doc, "Auth + roles working end-to-end")
    numbered(doc, "Core entities (venues / trainers / users / clubs) CRUD + dashboards")
    numbered(doc, "Booking availability + anti-double-book + hold expiry")
    numbered(doc, "Razorpay + status sync (venues, games, trainings, fees, fundraising)")
    numbered(doc, "Admin ops + venue verification + basic email")
    numbered(doc, "Staging smoke tests, monitoring, production cutover")

    heading(doc, "Band team sequence", 2)
    numbered(doc, "Auth + artist / client / venue roles")
    numbered(doc, "Profiles + categories / locations")
    numbered(doc, "Discovery (list + detail)")
    numbered(doc, "Booking request / accept / calendar integrity")
    numbered(doc, "Payments + booking confirmation")
    numbered(doc, "Messaging + notifications")
    numbered(doc, "Reviews + earnings")
    numbered(doc, "Staging smoke tests → production")

    heading(doc, "Practical daily usage", 2)
    bullet(doc, " Morning: pick ONE area (example: Sports fundraising, or Band bookings).")
    bullet(doc, " Use only that area’s skills from Section 5.")
    bullet(doc, " Definition of done: API + DB rule + UI + auth check + empty/error states.")
    bullet(doc, " Do not jump to analytics/polish before booking/payments work.")
    bullet(doc, " Cursor skill for Sports: mokijo-sports-only. Band teammates treat Sports paths as protected.")

    heading(doc, "Smoke tests before launch", 2)
    add_table(
        doc,
        ["Product", "Must-pass path"],
        [
            ["Sports", "Register → login → book venue (hold → pay → confirmed) → cancel/refund"],
            ["Sports", "Trainer publish → player enroll/pay → appears in registrations"],
            ["Sports", "Club create fundraising campaign → donate → raised total increases"],
            ["Sports", "Create match → go live → scoreboard updates"],
            ["Sports", "Super Admin approve club; Mukijo Admin verify venue"],
            ["Band", "Register → profile → search artist/venue → book → message → pay → review"],
        ],
        header_fill=C_HEADER,
        col_widths=[1.5, 5.8],
    )

    # ============================================================
    heading(doc, "9. Folder map (Sports vs Band — do not mix)", 1)

    heading(doc, "Mokijo Sports — permitted", 2)
    add_table(
        doc,
        ["Layer", "Paths"],
        [
            [
                "Frontend",
                "frontend/app/mokijo/, trainings/, scoreboard/, venue/, venues/, bookings/, user-dashboard/, super-admin/, components/auth/, dashboard/, trainer-dashboard/, venue-dashboard/, venue/, register/",
            ],
            [
                "Backend",
                "backend/app/api/mokijo/, activities/, events/, games/, matches/, trainer/, venue_owner/, venues/, shared/superadmin/, models/models.py, models/schemas.py, services/hold_expiry.py",
            ],
        ],
        header_fill=C_CLUB,
        col_widths=[1.4, 5.9],
    )

    heading(doc, "Band Connect — permitted (Sports team: do not touch)", 2)
    add_table(
        doc,
        ["Layer", "Paths"],
        [
            ["Frontend", "frontend/app/band/, frontend/app/artist/, frontend/components/band/, frontend/components/artist/"],
            [
                "Backend",
                "backend/app/api/band/, backend/app/api/band_*/, band_models.py, band_schemas.py, features/artists, categories, earnings, locations, reviews",
            ],
        ],
        header_fill=C_BAND,
        col_widths=[1.4, 5.9],
    )

    heading(doc, "Shared files — additive only", 2)
    para(
        doc,
        "backend/app/main.py, backend routers, frontend global CSS / app shell, package.json, requirements.txt: never delete or comment out the other module’s routes. Never drop the other module’s SQL tables (band_* vs sports/mokijo tables).",
    )

    callout(
        doc,
        "End of handbook",
        "Print or share this Word file with the team. For a new teammate: read Section 7.1 (who), 7.4 (fundraising example), then Section 5 (which skill where), then Section 9 (folders). Then pick one pin-to-pin flow and ship it end-to-end.",
    )

    doc.save(OUT)
    print(f"Wrote {OUT}")


if __name__ == "__main__":
    build()
